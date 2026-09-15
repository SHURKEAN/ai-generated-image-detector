"""Build the one-page DEIK.AI Challenge project description."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = (
    PROJECT_ROOT
    / "competition_submission"
    / "DEIK_AI_Challenge_2026_Project_Description_Compliant.docx"
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=80, bottom=65, end=80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def style_run(run, size: float = 12, bold: bool = False, color="000000") -> None:
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def remove_paragraph_borders(paragraph_or_style) -> None:
    element = getattr(paragraph_or_style, "_p", None)
    if element is not None:
        properties = element.get_or_add_pPr()
    else:
        properties = paragraph_or_style.element.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is not None:
        properties.remove(borders)


def add_paragraph(
    document: Document,
    text: str,
    *,
    heading: bool = False,
    before: float = 0,
    after: float = 4,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_with_next = heading
    style_run(paragraph.add_run(text), bold=heading)


def build_document() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.25)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(4)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(1)
    title.paragraph_format.keep_with_next = True
    style_run(title.add_run("AI Generated Image Detection Application"), 12, True)
    remove_paragraph_borders(document.styles["Title"])
    remove_paragraph_borders(title)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(7)
    subtitle.paragraph_format.keep_with_next = True
    style_run(
        subtitle.add_run("DEIK.AI Challenge 2026    Category 2A Open AI Project"),
        12,
        True,
    )

    add_paragraph(document, "Project summary", heading=True, after=2)
    add_paragraph(
        document,
        "This project provides a working application that estimates whether an uploaded image is a genuine photograph or AI-generated. The user receives a Real or AI-generated result, an AI probability, the selected detector version, and a warning that the result is probabilistic screening evidence rather than forensic proof. The application addresses a practical need in media verification, education, journalism, and digital-content review.",
    )

    add_paragraph(document, "Research foundation", heading=True, before=1, after=2)
    add_paragraph(
        document,
        "The prototype extends the published study by Zeyad Qasem, Ilham Gafarov, and Mokhaled N. A. Al-Hamadani, titled Hyperparameter-Aware Evaluation of Deep Learning Architectures for AI-Generated Image Detection. The study compared ResNet-50, EfficientNetV2-S, ViT-B/16, and Xception on the 120,000-image CIFAKE dataset. Its tuned weighted soft-voting ensemble achieved 99.07 percent accuracy and F1-score on the official 20,000-image test set. The four published checkpoints, preprocessing rules, class mapping, and ensemble weights remain unchanged and are protected by recorded SHA-256 hashes.",
    )

    add_paragraph(document, "Competition prototype", heading=True, before=1, after=2)
    add_paragraph(
        document,
        "Testing with an ordinary phone photograph exposed a domain shift: CIFAKE contains small 32 by 32 images and does not represent normal camera photographs. The competition version therefore uses declared input-domain routing. Native CIFAKE-scale inputs use the published ensemble, while normal-resolution images use the Community Forensics CVPR 2025 detector, a ViT-S/16 model developed for cross-generator detection. This prevents the application from presenting the published CIFAKE result as universal real-world accuracy. A separate CLIP ViT-L/14 head was trained on 5,931 registered research images, but retained as an experiment because it did not meet the promotion criteria.",
    )

    add_paragraph(document, "Evaluation results", heading=True, before=1, after=3)
    table = document.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(6.2), Cm(2.3), Cm(3.6), Cm(5.7))
    headers = ("Evaluation", "Images", "Result", "Observed errors")
    rows = (
        ("Published CIFAKE ensemble test", "20,000", "99.07% accuracy", "186 incorrect classifications"),
        ("Held-out real-world validation", "2,306", "99.83% balanced accuracy", "3 real false alarms; 1 AI miss"),
        ("Separate external smoke check", "6", "6 of 6 correct", "Small diagnostic set; not a benchmark"),
    )
    set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.width = widths[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if column_index in (0, 3)
                else WD_ALIGN_PARAGRAPH.CENTER
            )
            value = headers[column_index] if row_index == 0 else rows[row_index - 1][column_index]
            style_run(
                paragraph.add_run(value),
                12,
                row_index == 0,
                "FFFFFF" if row_index == 0 else "000000",
            )
            if row_index == 0:
                set_cell_shading(cell, "1F4E78")
            elif row_index % 2 == 1:
                set_cell_shading(cell, "F2F6FA")

    add_paragraph(
        document,
        "Original contribution and responsible use",
        heading=True,
        before=6,
        after=2,
    )
    add_paragraph(
        document,
        "The competition contribution converts the published experiment into a reproducible application with protected research checkpoints, model routing, source-separated evaluation, per-image provenance, licence tracking, duplicate detection, and a tested Gradio interface. Community Forensics and its dataset are cited academic resources. The application does not claim certainty and should not be used alone to accuse a person of creating or manipulating an image. Performance may change with unseen generators, screenshots, resizing, compression, or deliberate evasion.",
    )

    add_paragraph(document, "Demonstration", heading=True, before=1, after=2)
    add_paragraph(
        document,
        "The demo will analyze a genuine phone photograph and an official DALL-E example, display their probabilities and detector routes, and explain why independent real-world validation remains necessary even when a model performs strongly on its original benchmark.",
        after=0,
    )

    properties = document.core_properties
    properties.title = "AI Generated Image Detection Application"
    properties.subject = "DEIK.AI Challenge 2026 project description"
    properties.author = "Zeyad Qasem"
    properties.keywords = "AI-generated image detection, DEIK.AI Challenge 2026"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
